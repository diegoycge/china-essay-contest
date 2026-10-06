import random
from datetime import timedelta
from decimal import Decimal

from django.core import mail
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from portal.models import (
    AccessToken,
    Assignment,
    ContestSettings,
    EmailLog,
    Essay,
    JudgeAttestation,
    Prompt,
    Referral,
    User,
)
from portal.services import assignment as assigner
from portal.services import emails, exports, intake, scoring

from .factories import (
    LOREM,
    MediaMixin,
    make_admin,
    make_pdf,
    state_json,
    make_essay,
    make_judge,
    sheet_post,
    submit_sheet,
    upload,
)

PASSWORD = "correct horse battery"


@override_settings(EMAIL_CONFIGURED=True)
class AccountTests(TestCase):
    def test_invitation_sets_password_once(self):
        admin = make_admin()
        self.client.force_login(admin)
        response = self.client.post(
            reverse("manage_users"),
            {
                "full_name": "New Judge",
                "email": "New@Example.test",
                "role": "judge",
                "language": "en",
            },
        )
        self.assertRedirects(response, reverse("manage_users"))
        judge = User.objects.get(email="new@example.test")
        self.assertFalse(judge.has_usable_password())
        self.assertEqual(len(mail.outbox), 1)
        link = [word for word in mail.outbox[0].body.split() if "/invite/" in word][0]
        path = link.replace("http://localhost:8000", "")

        self.client.logout()
        self.assertContains(self.client.get(path), "New Judge")
        weak = self.client.post(path, {"password1": "short", "password2": "short"})
        self.assertEqual(weak.status_code, 200)
        done = self.client.post(path, {"password1": PASSWORD, "password2": PASSWORD})
        self.assertRedirects(done, reverse("home"), fetch_redirect_response=False)
        judge.refresh_from_db()
        self.assertTrue(judge.check_password(PASSWORD))
        # The link is spent.
        self.client.logout()
        self.assertEqual(self.client.get(path).status_code, 410)

    def test_only_token_hash_is_stored(self):
        judge = make_judge()
        _row, raw = AccessToken.issue(judge, AccessToken.PURPOSE_INVITE, timedelta(days=1))
        self.assertFalse(AccessToken.objects.filter(token_hash=raw).exists())
        self.assertIsNotNone(AccessToken.lookup(raw, AccessToken.PURPOSE_INVITE))
        self.assertIsNone(AccessToken.lookup(raw, AccessToken.PURPOSE_RESET))

    @override_settings(LOGIN_MAX_FAILURES=3)
    def test_sign_in_and_lockout(self):
        make_judge(email="j@example.test")
        url = reverse("login")
        for _ in range(3):
            self.assertContains(
                self.client.post(url, {"email": "j@example.test", "password": "wrong"}),
                "not correct",
            )
        locked = self.client.post(url, {"email": "j@example.test", "password": PASSWORD})
        self.assertContains(locked, "Too many failed attempts")
        self.assertNotIn("_auth_user_id", self.client.session)

    def test_deactivated_user_cannot_sign_in(self):
        make_judge(email="gone@example.test", is_active=False)
        response = self.client.post(
            reverse("login"), {"email": "gone@example.test", "password": PASSWORD}
        )
        self.assertContains(response, "not correct")

    def test_password_reset_does_not_reveal_accounts(self):
        make_judge(email="j@example.test")
        for address in ("j@example.test", "nobody@example.test"):
            response = self.client.post(reverse("forgot_password"), {"email": address})
            self.assertContains(response, "If that address belongs to an account")
        self.assertEqual(len(mail.outbox), 1)


class AccessControlTests(TestCase):
    def setUp(self):
        self.admin = make_admin()
        self.judge = make_judge("Judge One")
        self.other = make_judge("Judge Two")
        self.essay = make_essay()
        Assignment.objects.create(essay=self.essay, judge=self.judge)
        Assignment.objects.create(essay=self.essay, judge=self.other)

    def test_anonymous_users_are_sent_to_sign_in(self):
        for name in ("judge_dashboard", "manage_dashboard", "manage_essays", "manage_results"):
            response = self.client.get(reverse(name))
            self.assertEqual(response.status_code, 302)
            self.assertIn(reverse("login"), response["Location"])

    def test_judges_cannot_reach_any_committee_page(self):
        self.client.force_login(self.judge)
        pages = [
            reverse("manage_dashboard"),
            reverse("manage_essays"),
            reverse("manage_essay", args=[self.essay.pk]),
            reverse("manage_review", args=[self.essay.pk]),
            reverse("manage_original", args=[self.essay.pk]),
            reverse("manage_users"),
            reverse("manage_results"),
            reverse("manage_export", args=["xlsx-identities"]),
            reverse("manage_logs"),
        ]
        for page in pages:
            self.assertEqual(self.client.get(page).status_code, 403, page)

    def test_admins_do_not_judge(self):
        self.client.force_login(self.admin)
        self.assertEqual(
            self.client.get(reverse("judge_sheet", args=[self.essay.code])).status_code, 403
        )

    def test_judge_sees_only_own_assignments(self):
        unassigned = make_essay(original_filename="other.pdf")
        self.client.force_login(self.judge)
        self.assertEqual(
            self.client.get(reverse("judge_sheet", args=[unassigned.code])).status_code, 404
        )
        self.assertEqual(
            self.client.post(
                reverse("judge_autosave", args=[unassigned.code]), sheet_post()
            ).status_code,
            404,
        )

    def test_essay_page_leaks_nothing_identifying(self):
        submit_sheet(self.essay, self.other, scores=(9, 9, 9, 9))
        self.client.force_login(self.judge)
        for url in (reverse("judge_dashboard"), reverse("judge_sheet", args=[self.essay.code])):
            html = self.client.get(url).content.decode()
            for secret in (
                "Wei Zhang",
                "wei.zhang",
                "Riverside",
                "zhang_wei_essay",
                "Judge Two",
                "9.00",
            ):
                self.assertNotIn(secret, html, f"{secret!r} leaked on {url}")
        self.assertIn(self.essay.code, html)

    def test_unreleased_or_disqualified_essays_are_hidden(self):
        self.client.force_login(self.judge)
        url = reverse("judge_sheet", args=[self.essay.code])
        self.assertEqual(self.client.get(url).status_code, 200)
        for status in (Essay.STATUS_REVIEW, Essay.STATUS_DISQUALIFIED, Essay.STATUS_WITHDRAWN):
            Essay.objects.filter(pk=self.essay.pk).update(status=status)
            self.assertEqual(self.client.get(url).status_code, 404, status)

    def test_statement_must_be_signed_first(self):
        Prompt.objects.create(number=1, title="Chips")
        prompt = Prompt.objects.create(number=2, title="Taiwan")
        newcomer = make_judge("Judge New", onboarded=False)
        Assignment.objects.create(essay=make_essay(original_filename="n.pdf"), judge=newcomer)
        self.client.force_login(newcomer)
        self.assertRedirects(
            self.client.get(reverse("judge_dashboard")), reverse("judge_onboarding")
        )
        self.assertRedirects(
            self.client.get(reverse("judge_sheet", args=[self.essay.code])),
            reverse("judge_onboarding"),
        )
        incomplete = self.client.post(reverse("judge_onboarding"), {"signed_name": "Judge New"})
        self.assertEqual(incomplete.status_code, 200)
        self.client.post(
            reverse("judge_onboarding"),
            {
                "read_rubric": "on",
                "no_beliefs": "on",
                "conflict_prompts": [prompt.pk],
                "signed_name": "Judge New",
            },
        )
        attestation = JudgeAttestation.objects.get(user=newcomer)
        self.assertEqual(list(attestation.conflict_prompts.all()), [prompt])
        self.assertEqual(self.client.get(reverse("judge_dashboard")).status_code, 200)


@override_settings(EMAIL_CONFIGURED=True)
class ScoringSheetTests(TestCase):
    def setUp(self):
        self.judge = make_judge("Judge One")
        self.essay = make_essay()
        self.assignment = Assignment.objects.create(essay=self.essay, judge=self.judge)
        self.url = reverse("judge_sheet", args=[self.essay.code])
        self.client.force_login(self.judge)

    def test_draft_can_be_partial_and_autosaved(self):
        response = self.client.post(
            reverse("judge_autosave", args=[self.essay.code]),
            {"originality": "6.5", "originality_comment": "Early thoughts."},
        )
        self.assertEqual(response.json()["ok"], True)
        self.assignment.refresh_from_db()
        self.assertEqual(self.assignment.status, Assignment.STATUS_DRAFT)
        self.assertEqual(self.assignment.sheet.originality, Decimal("6.5"))
        self.assertIsNone(self.assignment.sheet.clarity)

    def test_submission_rules(self):
        cases = {
            "Score every criterion": sheet_post(depth=""),
            "Write a justification for every criterion": sheet_post(evidence_comment=" "),
            "at least 100 are required": sheet_post(
                originality_comment="Fine.",
                clarity_comment="Fine.",
                evidence_comment="Fine.",
                depth_comment="Fine.",
            ),
            "Explain the flag or referral": sheet_post(plagiarism_suspected="on"),
            "Tick the confirmation": sheet_post(confirmed=None),
        }
        for expected, data in cases.items():
            response = self.client.post(self.url, data)
            self.assertContains(response, expected)
            self.assignment.refresh_from_db()
            self.assertEqual(self.assignment.status, Assignment.STATUS_DRAFT, expected)

    def test_scores_must_be_half_point_steps_within_range(self):
        for bad in ("7.3", "10.5", "-1", "abc"):
            response = self.client.post(self.url, sheet_post(originality=bad))
            self.assertEqual(response.status_code, 200)
            self.assignment.refresh_from_db()
            self.assertNotEqual(self.assignment.status, Assignment.STATUS_SUBMITTED, bad)

    def test_submit_locks_the_sheet(self):
        self.assertRedirects(self.client.post(self.url, sheet_post()), reverse("judge_dashboard"))
        self.assignment.refresh_from_db()
        self.assertEqual(self.assignment.status, Assignment.STATUS_SUBMITTED)
        self.assertEqual(self.assignment.sheet.average, Decimal("7.125"))
        self.assertEqual(self.assignment.sheet.total, Decimal("28.5"))
        self.assertEqual(self.client.post(self.url, sheet_post(originality="1.0")).status_code, 403)
        self.assertEqual(
            self.client.post(
                reverse("judge_autosave", args=[self.essay.code]), sheet_post()
            ).status_code,
            409,
        )
        self.assignment.sheet.refresh_from_db()
        self.assertEqual(self.assignment.sheet.originality, Decimal("7.5"))

    def test_flags_create_referrals_for_the_chair(self):
        make_admin()
        self.client.post(
            self.url,
            sheet_post(
                refer_ai_markers="on",
                plagiarism_suspected="on",
                flag_notes="Citation 4 does not exist.",
            ),
        )
        kinds = set(Referral.objects.values_list("kind", flat=True))
        self.assertEqual(kinds, {Referral.KIND_AI, Referral.KIND_PLAGIARISM})
        self.assertTrue(any("referred" in message.subject for message in mail.outbox))

    def test_recusal_returns_essay_and_can_be_reassigned(self):
        admin = make_admin()
        replacement = make_judge("Judge Two")
        self.client.post(
            reverse("judge_recuse", args=[self.essay.code]), {"reason": "I have published on this."}
        )
        self.assignment.refresh_from_db()
        self.assertEqual(self.assignment.status, Assignment.STATUS_RECUSED)
        self.assertEqual(self.client.get(self.url).status_code, 404)
        referral = Referral.objects.get(kind=Referral.KIND_STANDPOINT)

        self.client.force_login(admin)
        self.client.post(
            reverse("manage_referrals"), {"referral": referral.pk, "action": "reassign"}
        )
        referral.refresh_from_db()
        self.assertFalse(referral.is_open)
        active = Assignment.objects.filter(essay=self.essay, status__in=Assignment.ACTIVE_STATUSES)
        self.assertEqual([a.judge for a in active], [replacement])

    def test_admin_can_reopen(self):
        self.client.post(self.url, sheet_post())
        self.client.force_login(make_admin())
        self.client.post(reverse("manage_sheet", args=[self.assignment.pk]), {"action": "reopen"})
        self.assignment.refresh_from_db()
        self.assertEqual(self.assignment.status, Assignment.STATUS_DRAFT)


class AssignmentTests(TestCase):
    def test_balanced_and_respects_conflicts_and_caps(self):
        taiwan = Prompt.objects.create(number=1, title="Taiwan")
        chips = Prompt.objects.create(number=2, title="Chips")
        judges = [make_judge(f"Judge {n}") for n in range(5)]
        judges[0].attestation.conflict_prompts.add(taiwan)
        judges[1].max_load = 2
        judges[1].save()
        essays = [
            make_essay(prompt=taiwan if n % 2 else chips, original_filename=f"{n}.pdf")
            for n in range(20)
        ]
        created, unfilled = assigner.auto_assign(rng=random.Random(7))
        self.assertEqual((len(created), unfilled), (40, []))
        for essay in essays:
            readers = list(essay.assignments.values_list("judge_id", flat=True))
            self.assertEqual(len(set(readers)), 2)
            if essay.prompt == taiwan:
                self.assertNotIn(judges[0].pk, readers)
        loads = {j.pk: j.assignments.count() for j in judges}
        self.assertEqual(loads[judges[1].pk], 2)
        # Judge 0 can only take the 10 "Chips" essays; the three unrestricted
        # judges share the remainder evenly.
        self.assertLessEqual(loads[judges[0].pk], 10)
        unrestricted = [loads[j.pk] for j in judges[2:]]
        self.assertLessEqual(max(unrestricted) - min(unrestricted), 1)
        # Running it again changes nothing.
        self.assertEqual(assigner.auto_assign(), ([], []))

    def test_reports_essays_it_cannot_fill_and_never_repeats_a_judge(self):
        only = make_judge("Only Judge")
        essay = make_essay()
        created, unfilled = assigner.auto_assign()
        self.assertEqual(len(created), 1)
        self.assertEqual([(e.pk, n) for e, n in unfilled], [(essay.pk, 1)])
        # A recused judge is never handed the same essay back.
        Assignment.objects.filter(judge=only).update(status=Assignment.STATUS_RECUSED)
        with self.assertRaises(assigner.NoEligibleJudge):
            assigner.assign(essay)

    def test_judges_wait_until_they_sign_the_statement(self):
        make_judge("Signed")
        unsigned = make_judge("Unsigned", onboarded=False)
        essay = make_essay()
        created, unfilled = assigner.auto_assign()
        self.assertEqual([a.judge.full_name for a in created], ["Signed"])
        self.assertEqual(len(unfilled), 1)
        with self.assertRaises(assigner.NoEligibleJudge):
            assigner.assign(essay, unsigned)

    def test_only_released_essays_are_assigned(self):
        make_judge("A")
        make_judge("B")
        make_essay(status=Essay.STATUS_REVIEW)
        self.assertEqual(assigner.auto_assign(), ([], []))


class ResultTests(TestCase):
    def setUp(self):
        self.config = ContestSettings.load()
        self.judges = [make_judge(f"Judge {n}") for n in range(4)]

    def test_divergence_and_rescore_cycle(self):
        essay = make_essay()
        submit_sheet(essay, self.judges[0], (8, 8, 8, 8))
        Assignment.objects.create(essay=essay, judge=self.judges[1])
        summary = scoring.summarize(essay, self.config)
        self.assertFalse(summary.divergent or summary.complete)

        submit_sheet(essay, self.judges[1], (6, 6, 6, 6))  # gap of exactly 2.0
        summary = scoring.summarize(essay, self.config)
        self.assertTrue(summary.divergent and summary.needs_rescore)
        self.assertFalse(summary.complete)

        assigner.assign(essay, self.judges[2], round_=Assignment.ROUND_RESCORE)
        summary = scoring.summarize(essay, self.config)
        self.assertFalse(summary.needs_rescore or summary.complete)

        submit_sheet(essay, self.judges[2], (6.5, 6.5, 6.5, 6.5), round_=Assignment.ROUND_RESCORE)
        summary = scoring.summarize(essay, self.config)
        self.assertTrue(summary.complete)
        self.assertEqual(summary.final, Decimal("6.83"))
        self.config.rescore_method = ContestSettings.RESCORE_CLOSEST_TWO
        self.assertEqual(scoring.summarize(essay, self.config).final, Decimal("6.25"))

    def test_gap_below_threshold_is_not_divergent(self):
        essay = make_essay()
        submit_sheet(essay, self.judges[0], (8, 8, 8, 7.5))
        submit_sheet(essay, self.judges[1], (6, 6, 6, 6.5))  # averages 7.875 and 6.125
        summary = scoring.summarize(essay, self.config)
        self.assertFalse(summary.divergent)
        self.assertTrue(summary.complete)
        self.assertEqual(summary.final, Decimal("7.00"))

    def test_removed_and_draft_sheets_do_not_count(self):
        essay = make_essay()
        submit_sheet(essay, self.judges[0], (8, 8, 8, 8))
        submit_sheet(essay, self.judges[1], (7, 7, 7, 7))
        removed = submit_sheet(essay, self.judges[2], (1, 1, 1, 1))
        removed.status = Assignment.STATUS_REMOVED
        removed.save()
        summary = scoring.summarize(essay, self.config)
        self.assertEqual(summary.final, Decimal("7.50"))
        self.assertFalse(summary.divergent)

    def test_ranking_is_per_division_with_shared_ranks(self):
        scores = {"hs": [(9, "a"), (7, "b"), (7, "c"), (5, "d")], "ug": [(6, "e")]}
        for division, rows in scores.items():
            for value, name in rows:
                essay = make_essay(division=division, original_filename=f"{name}.pdf")
                submit_sheet(essay, self.judges[0], (value,) * 4)
                submit_sheet(essay, self.judges[1], (value,) * 4)
        pending = make_essay(division="hs", original_filename="p.pdf")
        submit_sheet(pending, self.judges[0], (10, 10, 10, 10))
        Assignment.objects.create(essay=pending, judge=self.judges[1])

        ranked = scoring.rank(scoring.summarize_all(self.config))
        self.assertEqual([s.rank for s in ranked["hs"]], [1, 2, 2, 4])
        self.assertEqual([s.rank for s in ranked["ug"]], [1])
        self.assertNotIn(pending.pk, [s.essay.pk for s in ranked["hs"]])

    def test_exports(self):
        essay = make_essay()
        sheet = submit_sheet(essay, self.judges[0], (8, 8, 8, 8)).sheet
        sheet.overall_comment = '=HYPERLINK("http://evil.example")'
        sheet.save()
        submit_sheet(essay, self.judges[1], (7, 7, 7, 7))
        from io import BytesIO

        from openpyxl import load_workbook

        plain = load_workbook(BytesIO(exports.build_workbook()))
        self.assertEqual(len(plain.sheetnames), 3)
        for name in plain.sheetnames:
            for row in plain[name].iter_rows(values_only=True):
                self.assertNotIn("Wei Zhang", [str(cell) for cell in row])
                for cell in row:
                    self.assertFalse(str(cell).startswith("="), cell)
        full = load_workbook(BytesIO(exports.build_workbook(include_identities=True)))
        identities = list(full[full.sheetnames[-1]].iter_rows(values_only=True))
        self.assertIn("Wei Zhang", identities[1])
        self.assertIn(7.5, identities[1])
        csv_text = exports.build_csv("scores").decode("utf-8-sig")
        self.assertIn(essay.code, csv_text)
        self.assertIn("'=HYPERLINK", csv_text)


@override_settings(EMAIL_CONFIGURED=True)
class ReminderTests(TestCase):
    def test_each_scheduled_reminder_goes_out_once(self):
        config = ContestSettings.load()
        config.initial_deadline = timezone.now() + timedelta(days=2, hours=12)
        config.save()
        busy, idle = make_judge("Busy Judge"), make_judge("Idle Judge")
        Assignment.objects.create(essay=make_essay(), judge=busy)
        submit_sheet(make_essay(original_filename="b.pdf"), idle)

        self.assertEqual(emails.send_due_reminders(), 1)  # the 3-day reminder
        self.assertEqual(emails.send_due_reminders(), 0)
        self.assertEqual(mail.outbox[0].to, [busy.email])
        later = timezone.now() + timedelta(days=2)
        self.assertEqual(emails.send_due_reminders(now=later), 1)  # the 1-day reminder
        self.assertEqual(emails.send_due_reminders(now=later), 0)
        after = config.initial_deadline + timedelta(hours=1)
        self.assertEqual(emails.send_due_reminders(now=after), 1)  # one overdue notice
        self.assertEqual(emails.send_due_reminders(now=after), 0)
        self.assertEqual(EmailLog.objects.filter(kind="reminder").count(), 3)

    def test_emails_use_the_judges_language(self):
        judge = make_judge("张评委", email="zh@example.test", language="zh-hans")
        Assignment.objects.create(essay=make_essay(), judge=judge)
        emails.notify_new_assignments(list(judge.assignments.all()))
        self.assertIn("篇", mail.outbox[0].subject)


class ReviewTests(MediaMixin, TestCase):
    def setUp(self):
        super().setUp()
        self.admin = make_admin()
        self.client.force_login(self.admin)

    def _upload(self, name="zw.pdf", **options):
        options.setdefault("title_page", ["Chips and Power", "Wei Zhang"])
        options.setdefault("header", "Zhang")
        options.setdefault("author", "Wei Zhang")
        body = [LOREM * 3 + " Written by Wei Zhang.", LOREM * 3]
        response = self.client.post(
            reverse("manage_upload_file"),
            {"division": "ug", "prompt": "", "file": upload(name, make_pdf(body, **options))},
        )
        self.assertEqual(response.json()["outcome"], "created")
        return Essay.objects.get(original_filename=name)

    def test_upload_review_release(self):
        essay = self._upload()
        url = reverse("manage_review", args=[essay.pk])
        page = self.client.get(url)
        self.assertContains(page, "Pages left out (title page): 1")
        self.assertContains(page, reverse("manage_page", args=[essay.pk, 2]))
        image = self.client.get(reverse("manage_page", args=[essay.pk, 2]))
        self.assertEqual(image["Content-Type"], "image/png")
        self.client.cookies["django_language"] = "zh-hans"
        self.assertContains(self.client.get(url), "已排除的页面（标题页）：1")
        self.client.cookies["django_language"] = "en"

        # The reviewer switches the name's box off: release is refused.
        Essay.objects.filter(pk=essay.pk).update(author_name="Wei Zhang")
        essay.refresh_from_db()
        intake.run_auto_redaction(essay)
        essay.save()
        boxes = [dict(b, on=False) if b["kind"] == "name" else b for b in essay.redactions]
        self.client.post(url, {"action": "approve", "state": state_json(essay, redactions=boxes)})
        essay.refresh_from_db()
        self.assertEqual(essay.status, Essay.STATUS_REVIEW)
        self.assertIn("Wei Zhang", essay.released_text)
        self.assertEqual(essay.release_dir, "")

        # Switched back on (and one box drawn by hand): released.
        drawn = {"p": 1, "x": 0.2, "y": 0.6, "w": 0.3, "h": 0.04, "kind": "manual", "on": True}
        boxes = [dict(b, on=True) if b["kind"] == "name" else b for b in boxes] + [drawn]
        self.client.post(url, {"action": "approve", "state": state_json(essay, redactions=boxes)})
        essay.refresh_from_db()
        self.assertEqual(essay.status, Essay.STATUS_APPROVED)
        self.assertEqual(essay.approved_by, self.admin)
        self.assertEqual(essay.released_pages, essay.page_count - 1)
        self.assertIn(drawn, essay.redactions)
        preview = self.client.get(reverse("manage_preview", args=[essay.pk]))
        self.assertContains(preview, reverse("manage_released_page", args=[essay.pk, 1]))
        self.assertEqual(
            self.client.get(reverse("manage_released_page", args=[essay.pk, 1]))["Content-Type"],
            "image/png",
        )

    def test_judges_get_released_pages_only_for_their_own_essays(self):
        essay = self._upload()
        url = reverse("manage_review", args=[essay.pk])
        self.client.post(url, {"action": "approve", "state": state_json(essay)})
        essay.refresh_from_db()
        self.assertEqual(essay.status, Essay.STATUS_APPROVED)
        mine, other = make_judge("Judge One"), make_judge("Judge Two")
        Assignment.objects.create(essay=essay, judge=mine)

        self.client.force_login(mine)
        sheet = self.client.get(reverse("judge_sheet", args=[essay.code])).content.decode()
        first = reverse("judge_page", args=[essay.code, 1])
        self.assertIn(first, sheet)
        for secret in ("Wei Zhang", "zw.pdf", essay.work_dir, essay.release_dir, "/manage/"):
            self.assertNotIn(secret, sheet)
        response = self.client.get(first)
        self.assertEqual((response.status_code, response["Content-Type"]), (200, "image/png"))
        self.assertEqual(response["Cache-Control"], "no-store")
        # Past the last released page (the dropped title page is not addressable).
        self.assertEqual(
            self.client.get(
                reverse("judge_page", args=[essay.code, essay.released_pages + 1])
            ).status_code,
            404,
        )
        # No route to the unredacted renders or the original.
        for name in ("manage_page", "manage_released_page"):
            self.assertEqual(self.client.get(reverse(name, args=[essay.pk, 1])).status_code, 403)
        self.assertEqual(
            self.client.get(reverse("manage_original", args=[essay.pk])).status_code, 403
        )

        self.client.force_login(other)
        self.assertEqual(self.client.get(first).status_code, 404)
        self.client.logout()
        self.assertEqual(self.client.get(first).status_code, 302)

        # Pulled back: the page is gone for the assigned judge too.
        self.client.force_login(self.admin)
        self.client.post(url, {"action": "unrelease"})
        self.client.force_login(mine)
        self.assertEqual(self.client.get(first).status_code, 404)

    def test_original_file_download_is_logged(self):
        essay = self._upload("a.pdf")
        response = self.client.get(reverse("manage_original", args=[essay.pk]))
        self.assertTrue(b"".join(response.streaming_content).startswith(b"%PDF"))
        from portal.models import AuditLog

        self.assertTrue(
            AuditLog.objects.filter(action="original_downloaded", target=essay.code).exists()
        )

    def test_failed_file_can_be_retried(self):
        with self.settings(PDF_TIMEOUT_SECONDS=0.001):
            response = self.client.post(
                reverse("manage_upload_file"),
                {"division": "ug", "prompt": "", "file": upload("slow.pdf", make_pdf())},
            )
        self.assertEqual(response.json()["outcome"], "failed")
        essay = Essay.objects.get()
        url = reverse("manage_review", args=[essay.pk])
        self.assertContains(self.client.get(url), "took too long")
        self.client.post(url, {"action": "retry"})
        essay.refresh_from_db()
        self.assertEqual((essay.status, essay.extraction_note), (Essay.STATUS_REVIEW, ""))
        self.assertGreater(essay.page_count, 0)

    def test_every_committee_page_renders(self):
        judge, second = make_judge("Judge One"), make_judge("Judge Two")
        prompt = Prompt.objects.create(number=1, title="Chips", title_zh="芯片")
        essay = make_essay(prompt=prompt)
        low = submit_sheet(essay, judge, (4, 4, 4, 4))
        submit_sheet(essay, second, (9, 9, 9, 9))
        Referral.objects.create(assignment=low, kind=Referral.KIND_AI, notes="Fabricated citation.")
        make_essay(status=Essay.STATUS_REVIEW, original_filename="r.pdf")
        pages = [
            reverse("manage_dashboard"),
            reverse("manage_essays") + "?q=zhang&status=approved&division=ug",
            reverse("manage_upload"),
            reverse("manage_essay", args=[essay.pk]),
            reverse("manage_review", args=[essay.pk]),
            reverse("manage_preview", args=[essay.pk]),
            reverse("manage_users"),
            reverse("manage_user", args=[judge.pk]),
            reverse("manage_user", args=[self.admin.pk]),
            reverse("manage_assignments"),
            reverse("manage_divergence"),
            reverse("manage_referrals"),
            reverse("manage_sheet", args=[low.pk]),
            reverse("manage_results"),
            reverse("manage_prompts"),
            reverse("manage_prompt", args=[prompt.pk]),
            reverse("manage_settings"),
            reverse("manage_logs"),
            reverse("manage_logs") + "?tab=email",
            reverse("account"),
        ]
        for language in ("en", "zh-hans"):
            self.client.cookies["django_language"] = language
            for page in pages:
                self.assertEqual(self.client.get(page).status_code, 200, f"{language} {page}")
        self.assertContains(self.client.get(reverse("manage_divergence")), essay.code)

    def test_every_judge_page_renders(self):
        judge = make_judge("Judge One")
        essay = make_essay(prompt=Prompt.objects.create(number=1, title="Chips", text="Why chips?"))
        Assignment.objects.create(essay=essay, judge=judge)
        done = make_essay(original_filename="d.pdf")
        submit_sheet(done, judge)
        self.client.force_login(judge)
        for language in ("en", "zh-hans"):
            self.client.cookies["django_language"] = language
            for page in (
                reverse("judge_dashboard"),
                reverse("judge_rubric"),
                reverse("judge_sheet", args=[essay.code]),
                reverse("judge_sheet", args=[done.code]),
                reverse("account"),
            ):
                self.assertEqual(self.client.get(page).status_code, 200, f"{language} {page}")
