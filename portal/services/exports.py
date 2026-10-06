"""Spreadsheet exports for the committee."""

import csv
import io

from django.utils import timezone
from django.utils.translation import gettext as _
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..models import CRITERIA_KEYS, Assignment, ContestSettings, Essay, User
from ..rubric import criterion_name
from . import scoring
from .text import count_words


MAX_CELL_CHARS = 32000  # Excel's limit is 32,767 per cell


def _yes_no(value):
    return _("Yes") if value else _("No")


def _local(value):
    return timezone.localtime(value).strftime("%Y-%m-%d %H:%M") if value else ""


def _prompt(essay):
    return str(essay.prompt) if essay.prompt else ""


def _num(value):
    return float(value) if value is not None else None


def ranking_rows(summaries, config):
    readers = max([len(s.submitted) for s in summaries] + [config.readers_per_essay])
    header = [
        _("Division"),
        _("Rank"),
        _("Essay ID"),
        _("Prompt"),
        _("Final score"),
        _("Readers submitted"),
        _("Spread"),
        _("Divergent"),
        _("Re-scored"),
        _("Finalist"),
        _("Scoring complete"),
    ] + [_("Reader %(n)s average") % {"n": n} for n in range(1, readers + 1)]
    ranked = scoring.rank(summaries)
    rows = []
    for division in sorted(ranked):
        complete = ranked[division]
        incomplete = [s for s in summaries if s.essay.division == division and not s.complete]
        incomplete.sort(key=lambda s: s.essay.code)
        for summary in complete + incomplete:
            essay = summary.essay
            averages = [_num(scoring.quantize(a)) for a in summary.averages]
            rows.append(
                [
                    essay.get_division_display(),
                    summary.rank,
                    essay.code,
                    _prompt(essay),
                    _num(summary.final),
                    len(summary.submitted),
                    _num(summary.spread),
                    _yes_no(summary.divergent),
                    _yes_no(summary.has_rescore),
                    _yes_no(essay.is_finalist),
                    _yes_no(summary.complete),
                ]
                + averages
                + [None] * (readers - len(averages))
            )
    return header, rows


def score_rows():
    header = (
        [
            _("Essay ID"),
            _("Division"),
            _("Prompt"),
            _("Judge"),
            _("Judge email"),
            _("Round"),
            _("Status"),
        ]
        + [criterion_name(key) for key in CRITERIA_KEYS]
        + [_("Total (of 40)"), _("Average (of 10)")]
        + [
            _("%(criterion)s — justification") % {"criterion": criterion_name(k)}
            for k in CRITERIA_KEYS
        ]
        + [
            _("Overall comments"),
            _("Words written"),
            _("Suspected plagiarism"),
            _("Referred: AI markers"),
            _("Referred: other"),
            _("Notes on flags"),
            _("Signature"),
            _("Submitted at"),
        ]
    )
    rows = []
    assignments = (
        Assignment.objects.exclude(status=Assignment.STATUS_REMOVED)
        .select_related("essay", "essay__prompt", "judge", "sheet")
        .order_by("essay__division", "essay__code", "assigned_at")
    )
    for assignment in assignments:
        essay = assignment.essay
        sheet = getattr(assignment, "sheet", None)
        row = [
            essay.code,
            essay.get_division_display(),
            _prompt(essay),
            assignment.judge.full_name,
            assignment.judge.email,
            assignment.get_round_display(),
            assignment.get_status_display(),
        ]
        if sheet:
            row += [_num(getattr(sheet, key)) for key in CRITERIA_KEYS]
            row += [_num(sheet.total), _num(scoring.quantize(sheet.average))]
            row += [getattr(sheet, f"{key}_comment") for key in CRITERIA_KEYS]
            row += [
                sheet.overall_comment,
                sum(count_words(text) for text in sheet.comment_texts()),
                _yes_no(sheet.plagiarism_suspected),
                _yes_no(sheet.refer_ai_markers),
                _yes_no(sheet.refer_other),
                sheet.flag_notes,
                sheet.signature,
                _local(assignment.submitted_at),
            ]
        else:
            row += [None] * (len(header) - len(row))
            if assignment.closed_reason:
                row[-3] = assignment.closed_reason
        rows.append(row)
    return header, rows


def identity_rows(summaries):
    by_essay = {s.essay.id: s for s in summaries}
    header = [
        _("Essay ID"),
        _("Author"),
        _("Email"),
        _("School"),
        _("Author’s level"),
        _("Judged in"),
        _("Prompt"),
        _("Original file"),
        _("Status"),
        _("Final score"),
        _("Rank"),
        _("Finalist"),
    ]
    rows = []
    for essay in Essay.objects.select_related("prompt").order_by("division", "code"):
        summary = by_essay.get(essay.id)
        rows.append(
            [
                essay.code,
                essay.author_name,
                essay.author_email,
                essay.author_school,
                essay.get_author_level_display(),
                essay.get_division_display(),
                _prompt(essay),
                essay.original_filename,
                essay.get_status_display(),
                _num(summary.final) if summary else None,
                summary.rank if summary else None,
                _yes_no(essay.is_finalist),
            ]
        )
    return header, rows


def judge_rows():
    header = [
        _("Judge"),
        _("Email"),
        _("Affiliation"),
        _("Assigned"),
        _("Submitted"),
        _("Outstanding"),
        _("Statement signed"),
        _("Declared conflicts"),
    ]
    rows = []
    judges = User.objects.filter(role=User.ROLE_JUDGE).prefetch_related(
        "assignments", "attestation__conflict_prompts"
    )
    for judge in judges:
        active = [a for a in judge.assignments.all() if a.is_active]
        submitted = [a for a in active if a.status == Assignment.STATUS_SUBMITTED]
        attestation = getattr(judge, "attestation", None)
        conflicts = ""
        if attestation:
            conflicts = "; ".join(str(p) for p in attestation.conflict_prompts.all())
            if attestation.conflict_notes:
                conflicts = (conflicts + " — " if conflicts else "") + attestation.conflict_notes
        rows.append(
            [
                judge.full_name,
                judge.email,
                judge.affiliation,
                len(active),
                len(submitted),
                len(active) - len(submitted),
                _local(attestation.signed_at) if attestation else "",
                conflicts,
            ]
        )
    return header, rows


def _safe(cell):
    """Make judge- or author-supplied text safe to put in a spreadsheet.

    Control characters (a vertical tab pasted from Word is enough) make the
    XLSX writer raise, and a leading = + - @ would be run as a formula.
    """
    if not isinstance(cell, str):
        return cell
    cell = ILLEGAL_CHARACTERS_RE.sub("", cell)
    if len(cell) > MAX_CELL_CHARS:
        cell = cell[: MAX_CELL_CHARS - 1] + "…"
    if cell[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + cell
    return cell


HEADER_FILL = PatternFill("solid", fgColor="12284C")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _add_sheet(workbook, title, header, rows):
    sheet = workbook.create_sheet(title=title[:31])
    sheet.append(header)
    for row in rows:
        sheet.append([_safe(cell) for cell in row])
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    for index, title_cell in enumerate(header, 1):
        longest = max([len(str(title_cell))] + [len(str(r[index - 1] or "")) for r in rows[:200]])
        sheet.column_dimensions[get_column_letter(index)].width = min(max(10, longest + 2), 60)
    return sheet


def build_workbook(include_identities=False):
    config = ContestSettings.load()
    summaries = scoring.summarize_all(config)
    workbook = Workbook()
    workbook.remove(workbook.active)
    _add_sheet(workbook, _("Rankings"), *ranking_rows(summaries, config))
    _add_sheet(workbook, _("Scores and comments"), *score_rows())
    _add_sheet(workbook, _("Judges"), *judge_rows())
    if include_identities:
        _add_sheet(workbook, _("Author identities"), *identity_rows(summaries))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_csv(kind):
    if kind == "rankings":
        config = ContestSettings.load()
        header, rows = ranking_rows(scoring.summarize_all(config), config)
    else:
        header, rows = score_rows()
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(header)
    writer.writerows([["" if cell is None else _safe(cell) for cell in row] for row in rows])
    # The BOM makes Excel open UTF-8 (including Chinese comments) correctly.
    return ("﻿" + buffer.getvalue()).encode("utf-8")
